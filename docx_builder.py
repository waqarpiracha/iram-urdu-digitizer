"""
Iram Urdu Digitizer — Word (.docx) builder.

- Every paragraph is right-to-left (bidi), complex-script font set to
  "Noto Nastaliq Urdu" with graceful fallback if Iram's Word lacks it.
- Words the OCR was unsure about (low confidence) AND still present
  verbatim in her final text are highlighted YELLOW — "please verify".
  Words she corrected no longer match, so they lose the highlight:
  highlight == "the computer guessed this, you didn't change it".
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Complex-script font: pretty Nastaliq where installed, harmless fallback
# where not (Word substitutes automatically).
CS_FONT = "Noto Nastaliq Urdu"
BODY_SIZE_PT = 16
TITLE_SIZE_PT = 22

VERIFY_NOTE_URDU = (
    "نوٹ: پیلے رنگ سے نشان زد الفاظ کمپیوٹر کو واضح نہیں تھے۔ "
    "براہ کرم ان کی تصدیق کر لیں۔"
)


def _set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)


def _run_font(run, size_pt):
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:cs"), CS_FONT)          # complex-script (Urdu) font
    rPr.append(rFonts)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(int(size_pt * 2)))
    rPr.append(szCs)
    rtl = OxmlElement("w:rtl")
    rPr.append(rtl)


def _strip_punct(w):
    return w.strip("،۔؟!؛:()\"',.? ")


def _add_urdu_paragraph(doc, text, highlight_words, size_pt=BODY_SIZE_PT,
                        bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT):
    """One RTL paragraph; words in highlight_words get yellow highlight."""
    p = doc.add_paragraph()
    p.alignment = align
    _set_rtl(p)
    hl = {_strip_punct(w) for w in highlight_words}
    for word in text.split():
        run = p.add_run(word + " ")
        _run_font(run, size_pt)
        run.bold = bold
        if _strip_punct(word) in hl:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def build_docx(pages, output_path):
    """pages: list of dicts {title, text, low_conf_words}.
    low_conf_words: iterable of word strings to highlight if still present."""
    doc = Document()

    # Base style: RTL-friendly defaults
    style = doc.styles["Normal"]
    style.font.size = None  # per-run sizing instead

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_rtl(title)
    r = title.add_run("اردو نصاب")
    _run_font(r, TITLE_SIZE_PT)
    r.bold = True

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_rtl(note)
    rn = note.add_run(VERIFY_NOTE_URDU)
    _run_font(rn, 11)

    for i, page in enumerate(pages, 1):
        _add_urdu_paragraph(
            doc, page.get("title") or f"صفحہ {i}",
            set(), size_pt=18, bold=True)
        text = (page.get("text") or "").strip()
        hl = set(page.get("low_conf_words") or [])
        if text:
            for para_text in text.split("\n"):
                para_text = para_text.strip()
                if para_text:
                    _add_urdu_paragraph(doc, para_text, hl)
        else:
            _add_urdu_paragraph(
                doc, "[اس صفحے پر کوئی متن نہیں پڑھا جا سکا]",
                set(), size_pt=12)

    doc.save(output_path)
    return output_path
